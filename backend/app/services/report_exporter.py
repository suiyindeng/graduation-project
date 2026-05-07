import base64
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any
from uuid import uuid4

from docx import Document
from docx.shared import Inches

from app.core.config import settings
from app.models.dataset import Dataset, DatasetRow, ModelRun


def _add_key_value_table(document: Document, title: str, values: dict[str, Any]) -> None:
    document.add_heading(title, level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "内容"
    for key, value in values.items():
        cells = table.add_row().cells
        cells[0].text = str(key)
        cells[1].text = str(value)


def _add_chart_image(document: Document, title: str, image_base64: str) -> None:
    if "," in image_base64:
        image_base64 = image_base64.split(",", 1)[1]
    image_bytes = base64.b64decode(image_base64)
    with NamedTemporaryFile(delete=False, suffix=".png") as tmp:
        tmp.write(image_bytes)
        tmp_path = Path(tmp.name)
    document.add_paragraph(title)
    document.add_picture(str(tmp_path), width=Inches(6.2))
    tmp_path.unlink(missing_ok=True)


def export_word_report(
    dataset: Dataset,
    rows: list[DatasetRow],
    model_run: ModelRun | None,
    chart_images: list[dict[str, str]],
    notes: str | None = None,
) -> Path:
    """Generate a Word report containing charts and forecast results."""
    document = Document()
    document.add_heading("自动数据可视化分析报告", level=1)
    document.add_paragraph(f"数据文件：{dataset.filename}")
    document.add_paragraph(f"清洗后规模：{dataset.rows_count} 行，{dataset.columns_count} 列")

    profile = dataset.profile_json or {}
    _add_key_value_table(
        document,
        "一、数据清洗概览",
        {
            "原始行数": profile.get("original_rows"),
            "清洗后行数": profile.get("cleaned_rows"),
            "删除重复行": profile.get("duplicates_removed"),
            "数值字段": ", ".join(profile.get("numeric_columns", [])),
            "日期字段": ", ".join(profile.get("date_columns", [])),
            "分类字段": ", ".join(profile.get("categorical_columns", [])),
        },
    )

    document.add_heading("二、系统推荐分析", level=2)
    for item in dataset.recommendations_json or []:
        document.add_paragraph(f"{item.get('title')}：{item.get('reason')}")

    if chart_images:
        document.add_heading("三、可视化图表", level=2)
        for chart in chart_images:
            try:
                _add_chart_image(document, chart.get("title", "图表"), chart.get("image_base64", ""))
            except Exception:
                document.add_paragraph(f"{chart.get('title', '图表')}：图片写入失败")

    if model_run:
        _add_key_value_table(
            document,
            "四、行情预测结果",
            {
                "预测目标": model_run.target_column,
                "日期字段": model_run.date_column or "未使用",
                "算法": model_run.algorithm,
                "评估指标": model_run.metrics_json,
                "摘要": model_run.summary,
            },
        )
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "预测期"
        table.rows[0].cells[1].text = "预测值"
        for row in model_run.prediction_json:
            cells = table.add_row().cells
            cells[0].text = str(row.get("period"))
            cells[1].text = str(row.get("predicted_value"))

    document.add_heading("五、数据样例", level=2)
    preview = [row.content_json for row in rows[:8]]
    if preview:
        columns = list(preview[0].keys())[:6]
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for index, column in enumerate(columns):
            table.rows[0].cells[index].text = column
        for record in preview:
            cells = table.add_row().cells
            for index, column in enumerate(columns):
                cells[index].text = str(record.get(column, ""))

    if notes:
        document.add_heading("六、备注", level=2)
        document.add_paragraph(notes)

    filename = f"analysis-report-{uuid4().hex}.docx"
    report_path = settings.report_dir / filename
    document.save(report_path)
    return report_path
